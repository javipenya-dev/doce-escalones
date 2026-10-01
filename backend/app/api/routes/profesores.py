import io
import random
import pandas as pd
from fastapi import APIRouter, Depends, HTTPException, status, Query, UploadFile, File
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, or_
from typing import Optional

from app.db.database import get_db
from app.core.deps import get_current_admin
from app.models.models import Usuario, RolEnum
from app.schemas.schemas import UsuarioCreate, UsuarioUpdate, UsuarioOut
from app.core.security import hash_pin

router = APIRouter()


def _limpiar_tel(val):
    if val is None or pd.isna(val):
        return None
    if isinstance(val, float):
        return str(int(val)).strip()
    texto = str(val).strip()
    return texto[:-2] if texto.endswith(".0") else texto


@router.get("", response_model=list[UsuarioOut])
async def listar_profesores(
    nombre: Optional[str] = Query(None, description="Buscar por nombre o apellidos"),
    activo: Optional[bool] = Query(None, description="Filtrar por activo/baja"),
    incluir_admins: bool = Query(False, description="Si true, incluye también admins (para asignar clases)"),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_admin),
):
    """
    Lista usuarios con rol profesor. Filtros opcionales.
    """
    roles = [RolEnum.profesor]
    if incluir_admins:
        roles.append(RolEnum.admin)

    stmt = select(Usuario).where(Usuario.rol.in_(roles))

    if nombre:
        like = f"%{nombre}%"
        stmt = stmt.where(or_(
            Usuario.nombre.ilike(like),
            Usuario.apellidos.ilike(like),
        ))

    if activo is not None:
        stmt = stmt.where(Usuario.activo == activo)

    stmt = stmt.order_by(Usuario.apellidos, Usuario.nombre)
    result = await db.execute(stmt)
    return result.scalars().all()


@router.post("", response_model=UsuarioOut, status_code=status.HTTP_201_CREATED)
async def crear_profesor(
    data: UsuarioCreate,
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_admin),
):
    profesor = Usuario(
        nombre    = data.nombre,
        apellidos = data.apellidos,
        email     = data.email or None,
        telefono  = data.telefono or None,
        color     = data.color or '#F26419',
        pin       = hash_pin(data.pin),
        rol       = data.rol,           # ← cambiado: antes RolEnum.profesor
        activo    = True,
    )
    db.add(profesor)
    await db.flush()
    await db.refresh(profesor)
    return profesor


# ── DESCARGAR PLANTILLA PROFESORES ──────────────────────────────────────────

@router.get("/plantilla")
async def descargar_plantilla_profesores(
    _: Usuario = Depends(get_current_admin),
):
    """Genera y descarga un Excel plantilla para importar profesores."""
    from openpyxl.styles import Font, PatternFill

    df = pd.DataFrame([
        {
            "nombre": "Ana",
            "apellidos": "Profesora Ejemplo",
            "email": "ana@12escalones.com",
            "telefono": "600111222",
            "pin": "1234",
            "color": "#16A34A",
        },
        {
            "nombre": "Carlos",
            "apellidos": "López Martín",
            "email": "",
            "telefono": "600333444",
            "pin": "",
            "color": "",
        },
    ])

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Profesores")
        ws = writer.sheets["Profesores"]
        for i, col in enumerate(df.columns):
            ws.column_dimensions[chr(65 + i)].width = max(len(col) + 4, 18)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="F26419")

    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_profesores.xlsx"'},
    )


# ── IMPORTAR PROFESORES DESDE EXCEL ─────────────────────────────────────────

@router.post("/importar")
async def importar_profesores_excel(
    file: UploadFile = File(...),
    dry_run: bool = Query(False, description="Si true, solo valida y devuelve preview sin guardar"),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_admin),
):
    """
    Importa profesores desde un Excel (.xlsx/.xlsm).

    Columnas esperadas (case-insensitive):
      - nombre (obligatoria)
      - apellidos (obligatoria)
      - email
      - telefono
      - pin (4-6 dígitos; si vacío → se genera uno automáticamente)
      - color (hex #RRGGBB; si vacío → naranja marca)
    """
    if not (file.filename.endswith(".xlsx") or file.filename.endswith(".xlsm")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel (.xlsx o .xlsm)")

    try:
        contenido = await file.read()
        df = pd.read_excel(io.BytesIO(contenido))
        df.columns = [str(c).strip().lower() for c in df.columns]

        if "nombre" not in df.columns or "apellidos" not in df.columns:
            raise HTTPException(
                status_code=400,
                detail="El Excel debe contener las columnas 'nombre' y 'apellidos'",
            )

        validas = []
        errores = []
        duplicados = []
        creadas = []

        for idx, row in df.iterrows():
            fila_num = idx + 2

            nombre = str(row.get("nombre", "")).strip()
            apellidos = str(row.get("apellidos", "")).strip()
            if not nombre or nombre.lower() == "nan":
                errores.append({"fila": fila_num, "motivo": "Falta el nombre"})
                continue
            if not apellidos or apellidos.lower() == "nan":
                errores.append({"fila": fila_num, "motivo": "Faltan los apellidos"})
                continue

            email = (str(row["email"]).strip()
                     if "email" in df.columns and pd.notna(row["email"]) else None)
            telefono = (_limpiar_tel(row["telefono"])
                        if "telefono" in df.columns else None)

            # PIN: aceptar o auto-generar
            pin_raw = None
            if "pin" in df.columns and pd.notna(row["pin"]):
                pin_raw = str(row["pin"]).strip()
                if pin_raw.endswith(".0"):
                    pin_raw = pin_raw[:-2]
            if not pin_raw or not pin_raw.isdigit() or len(pin_raw) not in (4, 5, 6):
                pin_raw = f"{random.randint(1000, 9999)}"

            # Color: aceptar o default
            color = "#F26419"
            if "color" in df.columns and pd.notna(row["color"]):
                c = str(row["color"]).strip()
                if c.startswith("#") and len(c) == 7:
                    color = c

            # Duplicado: mismo nombre + apellidos
            existe_res = await db.execute(
                select(Usuario).where(
                    Usuario.nombre == nombre,
                    Usuario.apellidos == apellidos,
                    Usuario.rol == RolEnum.profesor,
                )
            )
            if existe_res.scalar_one_or_none():
                duplicados.append({
                    "fila": fila_num, "nombre": nombre, "apellidos": apellidos,
                    "motivo": "Ya existe un profesor con ese nombre y apellidos",
                })
                continue

            # Duplicado de email (si viene)
            if email:
                existe_email = await db.execute(
                    select(Usuario).where(Usuario.email == email)
                )
                if existe_email.scalar_one_or_none():
                    duplicados.append({
                        "fila": fila_num, "nombre": nombre, "apellidos": apellidos,
                        "motivo": f"El email {email} ya está en uso",
                    })
                    continue

            validas.append({
                "fila": fila_num,
                "nombre": nombre,
                "apellidos": apellidos,
                "email": email,
                "telefono": telefono,
                "pin": pin_raw,
                "color": color,
            })

            if not dry_run:
                nuevo = Usuario(
                    nombre=nombre,
                    apellidos=apellidos,
                    email=email,
                    telefono=telefono,
                    color=color,
                    pin=hash_pin(pin_raw),
                    rol=RolEnum.profesor,
                    activo=True,
                )
                db.add(nuevo)
                creadas.append({
                    "fila": fila_num,
                    "nombre": nombre,
                    "apellidos": apellidos,
                    "pin_temporal": pin_raw,
                })

        if not dry_run and creadas:
            await db.commit()

        return {
            "dry_run":        dry_run,
            "total_filas":    len(df),
            "num_validas":    len(validas),
            "num_errores":    len(errores),
            "num_duplicados": len(duplicados),
            "num_creados":    len(creadas),
            "validas":        validas[:50],
            "errores":        errores,
            "duplicados":     duplicados,
            "creados":        creadas,
            "mensaje": (
                f"Previsualización: {len(validas)} profesores listos para importar"
                if dry_run else
                f"Importación completada: {len(creadas)} profesores creados"
            ),
        }

    except HTTPException:
        raise
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=500, detail=f"Error al procesar el Excel: {str(e)}")


@router.put("/{profesor_id}", response_model=UsuarioOut)
async def actualizar_profesor(
    profesor_id: int,
    data: UsuarioUpdate,
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_admin),
):
    result = await db.execute(
        select(Usuario).where(Usuario.id == profesor_id)   # ← sin filtro de rol
    )
    profesor = result.scalar_one_or_none()
    if not profesor:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    update_data = data.model_dump(exclude_unset=True)

    if update_data.get("pin"):
        update_data["pin"] = hash_pin(update_data["pin"])
    else:
        update_data.pop("pin", None)

    if "email" in update_data and not update_data["email"]:
        update_data["email"] = None
    if "telefono" in update_data and not update_data["telefono"]:
        update_data["telefono"] = None

    for campo, valor in update_data.items():
        setattr(profesor, campo, valor)

    await db.flush()
    await db.refresh(profesor)
    return profesor


@router.delete("/{profesor_id}", status_code=status.HTTP_204_NO_CONTENT)
async def dar_baja_profesor(
    profesor_id: int,
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_admin),
):
    """Soft-delete: marca el usuario como inactivo (activo=False)."""
    result = await db.execute(
        select(Usuario).where(Usuario.id == profesor_id)   # ← sin filtro de rol
    )
    profesor = result.scalar_one_or_none()
    if not profesor:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    # Protección: no permitir quedarse sin admins activos
    if profesor.rol == RolEnum.admin:
        count_admin = await db.execute(
            select(Usuario).where(
                Usuario.rol == RolEnum.admin,
                Usuario.activo == True,
                Usuario.id != profesor_id,
            )
        )
        otros = count_admin.scalars().all()
        if not otros:
            raise HTTPException(
                status_code=400,
                detail="No puedes dar de baja al último administrador activo",
            )

    profesor.activo = False
    await db.flush()